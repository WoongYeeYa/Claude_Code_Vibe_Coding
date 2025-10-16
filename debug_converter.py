"""
DOCX to HWP 변환 디버깅 스크립트
"""
import sys
sys.path.insert(0, 'C:/workspace/Claude_Code_Vibe_Coding/backend')

from docx import Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
import win32com.client
import pythoncom

def debug_convert():
    """디버깅용 변환 함수"""

    # DOCX 파일 로드
    docx_path = 'C:/workspace/Claude_Code_Vibe_Coding/uploads/test_document.docx'
    doc = Document(docx_path)

    print("=== DOCX 파일 분석 ===")
    print(f"전체 문단 수: {len(doc.paragraphs)}")
    print(f"전체 표 수: {len(doc.tables)}\n")

    # 한글 초기화
    pythoncom.CoInitialize()
    hwp = win32com.client.Dispatch("HWPFrame.HwpObject")

    try:
        hwp.Visible = False
    except:
        pass

    hwp.RegisterModule("FilePathCheckDLL", "FilePathCheckerModule")

    # 새 문서 생성
    hwp.HAction.Run("FileNew")
    print("=== HWP 문서 생성 ===\n")

    # 문서 요소 순회
    element_count = 0
    for element in doc.element.body:
        element_count += 1

        if isinstance(element, CT_P):  # 문단
            # 해당 문단 찾기
            for para in doc.paragraphs:
                if para._element == element:
                    print(f"[문단 {element_count}] 스타일: {para.style.name}")
                    print(f"  텍스트: {repr(para.text)}")
                    print(f"  Runs 수: {len(para.runs)}")

                    # Run별 처리
                    text_inserted = False
                    for run_idx, run in enumerate(para.runs):
                        print(f"    Run {run_idx}: {repr(run.text)} (비어있음: {not run.text})")

                        if run.text:
                            # 텍스트 입력
                            hwp.HAction.GetDefault("InsertText", hwp.HParameterSet.HInsertText.HSet)
                            hwp.HParameterSet.HInsertText.Text = run.text
                            hwp.HAction.Execute("InsertText", hwp.HParameterSet.HInsertText.HSet)
                            text_inserted = True
                            print(f"      -> HWP에 입력됨")

                    if not text_inserted and para.text.strip():
                        print(f"  >> 대체 방법으로 입력: {repr(para.text)}")
                        hwp.HAction.GetDefault("InsertText", hwp.HParameterSet.HInsertText.HSet)
                        hwp.HParameterSet.HInsertText.Text = para.text
                        hwp.HAction.Execute("InsertText", hwp.HParameterSet.HInsertText.HSet)

                    # 문단 끝
                    hwp.HAction.Run("BreakPara")
                    print()
                    break

        elif isinstance(element, CT_Tbl):  # 표
            print(f"[표 {element_count}]")

            for table in doc.tables:
                if table._element == element:
                    rows = len(table.rows)
                    cols = len(table.columns)
                    print(f"  크기: {rows}행 x {cols}열")

                    # 표 생성
                    hwp.HAction.GetDefault("TableCreate", hwp.HParameterSet.HTableCreation.HSet)
                    hwp.HParameterSet.HTableCreation.Rows = rows
                    hwp.HParameterSet.HTableCreation.Cols = cols
                    hwp.HParameterSet.HTableCreation.WidthType = 2
                    hwp.HParameterSet.HTableCreation.HeightType = 0
                    hwp.HParameterSet.HTableCreation.CreateItemArray("ColWidth", cols)

                    col_width = int(60000 / cols)
                    for i in range(cols):
                        hwp.HParameterSet.HTableCreation.ColWidth.SetItem(i, col_width)

                    hwp.HAction.Execute("TableCreate", hwp.HParameterSet.HTableCreation.HSet)

                    # 첫 번째 셀로 이동
                    hwp.Run("TableCellBlock")
                    hwp.Run("Cancel")

                    # 셀 내용 채우기
                    for row_idx, row in enumerate(table.rows):
                        for col_idx, cell in enumerate(row.cells):
                            cell_text = cell.text.strip()
                            print(f"  셀 [{row_idx},{col_idx}]: {repr(cell_text)}")

                            if cell_text:
                                hwp.HAction.GetDefault("InsertText", hwp.HParameterSet.HInsertText.HSet)
                                hwp.HParameterSet.HInsertText.Text = cell_text
                                result = hwp.HAction.Execute("InsertText", hwp.HParameterSet.HInsertText.HSet)
                                print(f"    -> 입력 결과: {result}")

                            # 다음 셀로 이동
                            if not (row_idx == rows - 1 and col_idx == cols - 1):
                                move_result = hwp.Run("TableRightCell")
                                print(f"    -> 셀 이동 결과: {move_result}")

                    # 표 밖으로
                    hwp.HAction.Run("TableOut")
                    hwp.HAction.Run("BreakPara")
                    print()
                    break

    # 저장
    output_path = 'C:/workspace/Claude_Code_Vibe_Coding/outputs/debug_output.hwp'
    print(f"=== 저장 중: {output_path} ===")
    save_result = hwp.SaveAs(output_path, "HWP", "")
    print(f"저장 결과: {save_result}\n")

    # 종료
    hwp.Quit()
    pythoncom.CoUninitialize()

    print("=== 완료 ===")

if __name__ == "__main__":
    debug_convert()
